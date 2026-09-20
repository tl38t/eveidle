#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
军团 NPC 模板导入器（开发工具）
--------------------------------
读取「军团NPC自建填写模板」xlsx 的「填写模板」页，把每个 NPC 块转换成游戏数据：
  - 场景中文名  -> eventType
  - 技能中文名  -> skillId
  - 性格中文名  -> personalityId
  - 技能品质    -> skillGrade (A/B/C/D，缺省 A)  ← 主技能档位（玩家 NPC 约定默认 A）
  - 副技能      -> 随机抽一个 ≠ 主技能的技能，固定 B 级（招募即得，参与数值贡献）
导出 js/data/legion/self-built-npcs.js（双环境：window.LEGION_NPC_SELF_BUILT / module.exports）。
注意：生成的 NPC 经 legion-npc.js 的 injectSelfBuiltIntoCandidates 注入「招募池」（不预招募），
      玩家花 ISK/LP 招募后才进军团；已解雇则记 dismissedSelfBuilt 不再复活。

用法：
  python tools/import-legion-npcs.py <模板.xlsx> <输出.js> [--repo D:/EVE-IDLE/EVEIDLE-WORKBUDDY-FRESH]
不传参数时用仓库内默认路径。

注意：本工具只转换「填写模板」页里有「名称」行的块；模板里的「说明/示例」行自动忽略。
内置随机角色台词库（DIALOGUE）完全不被触碰——自建 NPC 的台词只写进 customDialogue 字段。
"""
import re
import sys
import os
import random

# ---------- 稳定映射（与 js/data/legion/* 真值逐字一致）----------
SKILL_NAME_TO_ID = {
    "矿脉勘探": "mining",
    "行星统筹": "planetaryIndustry",
    "熔炉调谐": "refining",
    "气云析取": "gasHarvesting",
    "舰构工程": "shipEngineering",
    "装备装配": "equipmentEngineering",
    "配方调制": "boosterEngineering",
    "激光火控": "laserOps",
    "炮台校准": "cannonOps",
    "制导算法": "missileOperations",
    "护盾整流": "shieldOperation",
    "装甲整备": "armorReinforcement",
    "船体加固": "hullEngineering",
    "电容节流": "capacitorManagement",
    "战利品搜寻": "lootSearch",
    "遗迹解析": "archaeologySpeed",
    "遗物鉴定": "archaeologyLoot",
    "产线调度": "autolineSpeed",
    "舰材回收": "shipComponentCostReduce",
    "训练教范": "xpGain",
    "薪资统筹": "wageReduce",
}
ALL_SKILL_IDS = list(SKILL_NAME_TO_ID.values())  # 副技能随机池（排除主技能后抽取）

PERSONALITY_NAME_TO_ID = {
    "冷静": "calm", "热情": "warm", "寡言": "taciturn", "傲慢": "arrogant",
    "慵懒": "lazy", "认真": "serious", "乐观": "optimistic", "悲观": "pessimistic",
    "神秘": "mystic", "机敏": "sharp", "直率": "blunt", "腹黑": "scheming",
}
SCENE_NAME_TO_EVENT = {
    "入职报到": "recruit",
    "每日问候": "dailyGreeting",
    "发薪到账": "salaryPaid",
    "欠薪待命": "salaryOverdue",
    "配舰完成": "shipAssigned",
    "换舰": "shipReplaced",
    "暂未配舰": "noShip",
    "错舰减半": "incompatibleShip",
    "升级": "levelUp",
    "技能里程碑": "skillMilestone",
    "满级": "maxLevel",
    "解雇": "dismiss",
}
GRADES = {"A", "B", "C", "D"}
SCENE_RE = re.compile(r"^\s*\d+\.\s*(.+?)\s*$")


def parse_blocks(ws):
    blocks = []
    max_row = ws.max_row
    r = 1
    while r <= max_row:
        a = ws.cell(row=r, column=1).value
        if a == "名称":
            name = ws.cell(row=r, column=2).value
            if not name:
                r += 1
                continue
            block = {"name": str(name).strip(), "skill": None, "grade": "A",
                     "personality": None, "customDialogue": {}}
            # 技能品质写在「名称」行的 E/F 列
            if ws.cell(row=r, column=5).value == "技能品质":
                g = ws.cell(row=r, column=6).value
                if g and str(g).strip().upper() in GRADES:
                    block["grade"] = str(g).strip().upper()
            # 扫描本块后续行（上限 25 行，到下一个「名称」或表尾为止）
            for rr in range(r + 1, min(r + 26, max_row + 1)):
                ca = ws.cell(row=rr, column=1).value
                if ca == "名称":
                    break
                if ca == "专精技能":
                    block["skill"] = ws.cell(row=rr, column=2).value
                    # 性格写在「专精技能」行的 E/F 列
                    if ws.cell(row=rr, column=5).value == "性格":
                        p = ws.cell(row=rr, column=6).value
                        if p:
                            block["personality"] = str(p).strip()
                elif ca is not None:
                    m = SCENE_RE.match(str(ca))
                    if m:
                        scene = m.group(1)
                        ev = SCENE_NAME_TO_EVENT.get(scene)
                        if ev:
                            lines = [ws.cell(row=rr, column=c).value for c in (2, 3, 4)]
                            lines = [str(x).strip() for x in lines if x and str(x).strip()]
                            if lines:
                                block["customDialogue"][ev] = lines
            blocks.append(block)
            r = rr_end(ws, r, max_row)
        else:
            r += 1
    return blocks


def rr_end(ws, start, max_row):
    # 找到下一个「名称」行，返回其行号；否则表尾+1
    for rr in range(start + 1, max_row + 2):
        if rr > max_row:
            return max_row + 1
        if ws.cell(row=rr, column=1).value == "名称":
            return rr
    return max_row + 1


def build_npcs(blocks):
    npcs = []
    seen_ids = {}
    warnings = []
    for b in blocks:
        skill_id = SKILL_NAME_TO_ID.get(b["skill"]) if b["skill"] else None
        if not skill_id:
            warnings.append("技能未识别，跳过块 name=%s skill=%r" % (b["name"], b["skill"]))
            continue
        pid = PERSONALITY_NAME_TO_ID.get(b["personality"]) if b["personality"] else None
        if b["personality"] and not pid:
            warnings.append("性格未识别 name=%s personality=%r（将留空，运行时随机）" % (b["name"], b["personality"]))
        # 副技能：随机抽一个 ≠ 主技能的技能，固定 B 级（玩家 NPC 约定：随机 B 级副技能）
        sec_choices = [sid for sid in ALL_SKILL_IDS if sid != skill_id]
        secondary_skill_id = random.choice(sec_choices)
        secondary_grade = "B"
        base = "self_" + re.sub(r"[^\w一-鿿]", "_", b["name"])
        npc_id = base
        if npc_id in seen_ids:
            seen_ids[npc_id] += 1
            npc_id = "%s_%d" % (base, seen_ids[npc_id])
        else:
            seen_ids[npc_id] = 0
        npcs.append({
            "npcId": npc_id,
            "name": b["name"],
            "personalityId": pid,
            "skillId": skill_id,
            "skillGrade": b["grade"],
            "secondarySkillId": secondary_skill_id,
            "secondarySkillGrade": secondary_grade,
            "customDialogue": b["customDialogue"],
        })
    return npcs, warnings


def emit_js(npcs, out_path):
    lines = []
    lines.append("// ================================================================")
    lines.append("// 军团 DLC —— 自建 NPC 数据（AUTO-GENERATED，勿手改）")
    lines.append("// 由 tools/import-legion-npcs.py 从「军团NPC自建填写模板」生成。")
    lines.append("// 这些 NPC 的台词只存在于 customDialogue 字段，绝不写入内置随机角色台词库。")
    lines.append("// 兼容双环境：浏览器挂 window.LEGION_NPC_SELF_BUILT，Node 下 module.exports。")
    lines.append("// ================================================================")
    lines.append("(function (root, factory) {")
    lines.append("  const mod = factory();")
    lines.append("  if (typeof module !== \"undefined\" && module.exports) module.exports = mod;")
    lines.append("  if (typeof window !== \"undefined\") window.LEGION_NPC_SELF_BUILT = mod;")
    lines.append("  else if (root) root.LEGION_NPC_SELF_BUILT = mod;")
    lines.append("})(typeof self !== \"undefined\" ? self : this, function () {")
    lines.append("  const SELF_BUILT_NPCS = [")
    for i, n in enumerate(npcs):
        cd = "{\n" + ",\n".join(
            "          %r: %r" % (k, v) for k, v in n["customDialogue"].items()
        ) + "\n        }" if n["customDialogue"] else "null"
        obj = (
            "    {\n"
            "      npcId: %r,\n"
            "      name: %r,\n"
            "      personalityId: %r,\n"
            "      skillId: %r,\n"
            "      skillGrade: %r,\n"
            "      secondarySkillId: %r,\n"
            "      secondarySkillGrade: %r,\n"
            "      customDialogue: %s\n"
            "    }" % (n["npcId"], n["name"], n["personalityId"], n["skillId"], n["skillGrade"], n["secondarySkillId"], n["secondarySkillGrade"], cd)
        )
        lines.append(obj + ("," if i < len(npcs) - 1 else ""))
    lines.append("  ];")
    lines.append("  return { SELF_BUILT_NPCS: SELF_BUILT_NPCS };")
    lines.append("});")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


def main():
    repo = "D:/EVE-IDLE/EVEIDLE-WORKBUDDY-FRESH"
    args = sys.argv[1:]
    xlsx = None
    out = None
    i = 0
    while i < len(args):
        if args[i] == "--repo":
            repo = args[i + 1]
            i += 2
        elif xlsx is None:
            xlsx = args[i]
            i += 1
        elif out is None:
            out = args[i]
            i += 1
        else:
            i += 1
    if xlsx is None:
        xlsx = "D:/Downloads/军团NPC自建填写模板(1).xlsx"
    if out is None:
        out = os.path.join(repo, "js/data/legion/self-built-npcs.js")

    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["填写模板"]
    blocks = parse_blocks(ws)
    npcs, warnings = build_npcs(blocks)
    emit_js(npcs, out)
    print("输入:", xlsx)
    print("输出:", out)
    print("解析到 NPC 块:", len(blocks), "生成:", len(npcs))
    for w in warnings:
        print("  [WARN]", w)
    for n in npcs:
        print("  +", n["npcId"], n["name"], n["skillId"], n["skillGrade"],
              "personality=", n["personalityId"], "scenes=", len(n["customDialogue"]))


if __name__ == "__main__":
    main()
